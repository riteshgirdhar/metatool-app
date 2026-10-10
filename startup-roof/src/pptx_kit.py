"""Small drawing kit on top of python-pptx so every visual stays an editable shape."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY = "10213B"
BLUE = "1765D1"
PALE = "EDF4FE"
WHITE = "FFFFFF"
TEXT = "253650"
MUTED = "61738A"
LINE = "D5E2F3"
MID = "4F8BE0"      # lighter blue for tints on navy
DEEP = "1A3358"     # navy tint for panels on navy
FONT = "Inter"

W, H = 13.333, 7.5
ML = 0.65  # left margin


def rgb(h):
    return RGBColor.from_string(h)


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def blank(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bgfill = s.background.fill
    bgfill.solid()
    bgfill.fore_color.rgb = rgb(bg)
    return s


def strip_style(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None, name=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if radius is not None and len(sh.adjustments) > 0:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    strip_style(sh)
    if name:
        sh.name = name
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return sh


def _apply_run(run, size, color, bold, italic=False, font=FONT, spacing=None):
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if spacing is not None:
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(spacing * 100)))


def fill_tf(tf, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, spacing=None, line=1.15, after=0, font=FONT, italic=False):
    """paras: list of paragraphs. Each paragraph is a str or a list of runs.
    A run is a str or (str, dict(size,color,bold,italic,spacing))."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        popts = {}
        if isinstance(p, dict):
            popts = p
            p = p["runs"]
        para.alignment = popts.get("align", align)
        para.line_spacing = popts.get("line", line)
        para.space_after = Pt(popts.get("after", after))
        if "before" in popts:
            para.space_before = Pt(popts["before"])
        runs = [p] if isinstance(p, (str, tuple)) else p
        for r in runs:
            if isinstance(r, str):
                r = (r, {})
            t, o = r
            run = para.add_run()
            run.text = t
            _apply_run(run, o.get("size", popts.get("size", size)), o.get("color", popts.get("color", color)),
                       o.get("bold", popts.get("bold", bold)), o.get("italic", italic), font,
                       o.get("spacing", popts.get("spacing", spacing)))
    return tf


def text(s, x, y, w, h, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=1.15, after=0, margin=0.0, name=None, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.auto_size = None
    fill_tf(tf, paras, size, color, bold, align, anchor, spacing, line, after, italic=italic)
    if name:
        tb.name = name
    return tb


def shape_text(sh, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, line=1.1, after=0, margin=0.08, spacing=None):
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    fill_tf(tf, paras, size, color, bold, align, anchor, spacing, line, after)
    return sh


def line(s, x1, y1, x2, y2, color=LINE, w=1.0, arrow=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    strip_style(c)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if arrow:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def polyline(s, pts, color=WHITE, w=3.0, closed=False, fill=None):
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(a), Inches(b)) for a, b in pts[1:]], close=closed)
    sh = fb.convert_to_shape()
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if w:
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    strip_style(sh)
    return sh


def wordmark(s, x, y, size=12, dark=False, align=PP_ALIGN.LEFT, w=3.0):
    c1 = WHITE if dark else NAVY
    c2 = MID if dark else BLUE
    return text(s, x, y, w, size / 72 * 1.4,
                [[("STARTUP", {"color": c1, "bold": True, "spacing": size * 0.12}),
                  (" ROOF", {"color": c2, "bold": True, "spacing": size * 0.12})]],
                size=size, align=align)


def pill(s, x, y, label, fill=PALE, color=BLUE, size=9, w=None, line_color=None):
    w = w or (0.3 + len(label) * size * 0.0098)
    sh = rect(s, x, y, w, 0.28, fill=fill, line=line_color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    shape_text(sh, [[(label, {"bold": True, "spacing": 0.8})]], size=size, color=color, margin=0.06)
    sh.text_frame.word_wrap = False
    return sh


def header(s, num, section, title, dark=False, title_size=28, tag=None, title_w=11.6):
    ec = MID if dark else BLUE
    tc = WHITE if dark else NAVY
    text(s, ML, 0.48, 8, 0.3, [[(f"{num:02d}", {"bold": True}), ("   " + section.upper(), {"bold": True})]],
         size=10, color=ec, spacing=1.5)
    text(s, ML, 0.82, title_w, 1.1, title, size=title_size, color=tc, bold=True, line=1.02)
    if tag:
        tag_w = 0.3 + len(tag) * 9 * 0.0098
        pill(s, W - ML - tag_w, 0.44, tag, fill=(DEEP if dark else PALE), color=(WHITE if dark else BLUE))


def footer(s, n, dark=False):
    c = "8EA0B8" if dark else MUTED
    wordmark(s, ML, 7.02, size=8.5, dark=dark)
    text(s, W - ML - 3, 7.02, 3, 0.2, [[("startuproof.com", {}), (f"     {n:02d}", {"bold": True})]],
         size=8.5, color=c, align=PP_ALIGN.RIGHT)


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body


def table(s, x, y, w, h, data, col_w, header_fill=NAVY, header_color=WHITE, size=11.5, header_size=11,
          row_fills=(WHITE, PALE), first_col_bold=True, first_col_color=NAVY, row_h=None):
    rows, cols = len(data), len(data[0])
    gt = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    # strip default style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(rows):
        if row_h:
            tbl.rows[r].height = Inches(row_h[r] if isinstance(row_h, (list, tuple)) else row_h)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            else:
                cell.fill.fore_color.rgb = rgb(row_fills[(r - 1) % len(row_fills)])
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = data[r][c]
            if r == 0:
                fill_tf(cell.text_frame, [[(val, {"bold": True, "spacing": 0.6})]], size=header_size, color=header_color)
            elif c == 0 and first_col_bold:
                fill_tf(cell.text_frame, val, size=size, color=first_col_color, bold=True)
            else:
                fill_tf(cell.text_frame, val, size=size, color=TEXT)
    return gt
