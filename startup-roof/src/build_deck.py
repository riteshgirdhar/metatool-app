"""Builds Startup_Roof_General_Presentation.pptx (20 slides, 16:9, editable shapes, presenter notes)."""
import io
import sys
from pathlib import Path

import qrcode
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from pptx_kit import *  # noqa: F401,F403

URL = "https://startuproof.com"

prs = None  # set by the builder before slides are drawn
B = lambda t, **o: (t, dict(bold=True, **o))  # noqa: E731


def qr_png(url, dark="10213B"):
    q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=20, border=2)
    q.add_data(url)
    q.make(fit=True)
    img = q.make_image(fill_color="#" + dark, back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def s_cover():
    global s
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.18, H, fill=BLUE)
    wordmark(s, ML + 0.1, 0.62, size=13, dark=True, w=4)
    text(s, ML + 0.1, 1.95, 7, 0.3, "A FOUNDER ECOSYSTEM  ·  UAE FIRST, GLOBAL AMBITION", size=10.5, color=MID, bold=True, spacing=1.6)
    text(s, ML + 0.1, 2.35, 7.2, 1.9, [[("STARTUP", {"color": WHITE}), ("\nROOF", {"color": WHITE})]], size=66, bold=True, line=0.9, spacing=2)
    text(s, ML + 0.1, 4.5, 6.4, 1.0, "A year-round ecosystem to help founders build, learn, connect and grow.", size=20, color="D6E2F3", line=1.15)
    rect(s, ML + 0.1, 5.65, 0.9, 0.04, fill=BLUE)
    text(s, ML + 0.1, 5.85, 6.6, 0.7,
         [[B("Home of Founder OS", color=WHITE), ("  —  the digital workspace for the founder journey.", {"color": "B7C6DC"})]],
         size=13.5)
    text(s, ML + 0.1, 6.75, 4, 0.3, "startuproof.com", size=12, color=MID, bold=True, spacing=0.8)

    # roof + four pillars visual (editable shapes)
    bx, by, bw = 7.75, 1.25, 4.95
    peak = (bx + bw / 2, by)
    base_y = by + 1.55
    polyline(s, [(bx - 0.25, base_y), peak, (bx + bw + 0.25, base_y)], color=WHITE, w=5)
    pill_labels = ["BUILD", "LEARN", "CONNECT", "GROW"]
    cw, gap = 0.98, (bw - 4 * 0.98) / 3
    floor = 6.15
    top = base_y + 0.3
    rect(s, bx - 0.05, top - 0.12, bw + 0.1, 0.05, fill="7F95B5")
    for i, lab in enumerate(pill_labels):
        x = bx + i * (cw + gap)
        rect(s, x, top, cw, floor - top, fill=[DEEP, "1C3D6B", "1F4A85", BLUE][i])
        text(s, x, floor + 0.15, cw, 0.3, lab, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, spacing=1.2)
    rect(s, bx - 0.25, floor, bw + 0.5, 0.03, fill="7F95B5")
    # founder at the centre, under the roof
    fo = rect(s, peak[0] - 0.55, base_y - 0.62, 1.1, 0.42, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    shape_text(fo, "FOUNDER", size=10, color=NAVY, bold=True, spacing=1.2)
    notes(s, "Open with the one-line idea: Startup Roof is a year-round ecosystem around the founder — guidance, tools, people and opportunities in one place. "
             "Founder OS is the digital workspace inside it. The visual is the idea in miniature: one roof over four things founders need — build, learn, connect, grow. "
             "Set expectations early: we are at an exploratory, pilot-shaping stage and this deck explains the concept and invites participation.")


def s_problem():
    global s
    s = blank(prs)
    header(s, 2, "The founder experience today", "Founders need more than information.\nThey need a practical path forward.", title_size=30)
    items = [
        ("Information is scattered", "Guidance sits across many websites, groups and programmes — hard to compare, harder to act on."),
        ("Tools do not connect", "Plans, tasks, documents and contacts live in separate apps with no shared picture of progress."),
        ("Advice lacks a next step", "General advice rarely translates into “what should I do this week?”"),
        ("The right people are hard to find", "Relevant mentors and specialists are often reached by chance, not by fit."),
        ("Introductions go nowhere", "Connections are made, but few lead to a conversation with a useful outcome."),
        ("Time and money are limited", "Founders must choose carefully where to spend scarce hours and budget."),
    ]
    y0 = 2.35
    for i, (h1, d) in enumerate(items):
        col, row = i % 2, i // 2
        x = ML + col * 4.15
        y = y0 + row * 1.42
        text(s, x, y, 0.5, 0.4, f"{i + 1:02d}", size=13, color=BLUE, bold=True)
        text(s, x + 0.55, y, 3.4, 0.35, h1, size=14.5, color=NAVY, bold=True)
        text(s, x + 0.55, y + 0.38, 3.35, 0.9, d, size=11.5, color=MUTED, line=1.2)
        line(s, x, y - 0.14, x + 3.9, y - 0.14, color=LINE, w=0.75)
    p = rect(s, 9.15, 2.2, 3.55, 4.45, fill=NAVY)
    text(s, 9.45, 2.48, 3.0, 0.3, "WHAT WE WILL INVESTIGATE", size=9.5, color=MID, bold=True, spacing=1.2)
    text(s, 9.45, 2.88, 3.0, 2.2,
         "Our working view is that founders move faster when guidance, tools and people are connected to a clear next step.",
         size=15.5, color=WHITE, bold=True, line=1.2)
    text(s, 9.45, 5.05, 3.0, 1.4,
         "The pilot will test this directly with founders — through interviews, sessions and follow-up — before we scale anything.",
         size=11.5, color="C9D6E8", line=1.25)
    footer(s, 2)
    notes(s, "These six pain points are the starting hypotheses, not research findings. They come from the experience of the people shaping the project and will be tested in founder interviews. "
             "The key message: the gap is rarely information — it is turning information into a practical next step with the right people around you.")


def s_what():
    global s
    s = blank(prs)
    header(s, 3, "What Startup Roof is", "One ecosystem around the founder.")
    cx, cy, R = 3.75, 4.45, 2.15
    outer = rect(s, cx - R, cy - R, 2 * R, 2 * R, fill=PALE, shape=MSO_SHAPE.OVAL)
    mid = rect(s, cx - 1.38, cy - 1.38, 2.76, 2.76, fill=None, line="B8CFEE", lw=1.0, shape=MSO_SHAPE.OVAL)
    core = rect(s, cx - 0.72, cy - 0.72, 1.44, 1.44, fill=NAVY, shape=MSO_SHAPE.OVAL)
    shape_text(core, [[B("Founder")]], size=15, color=WHITE)
    nodes = [("Community", cx, cy - 1.38), ("Programmes", cx + 1.38, cy), ("Services", cx, cy + 1.38), ("Founder OS", cx - 1.38, cy)]
    for lab, nx, ny in nodes:
        is_os = lab == "Founder OS"
        n = rect(s, nx - 0.72, ny - 0.24, 1.44, 0.48, fill=(BLUE if is_os else WHITE), line=(None if is_os else "B8CFEE"),
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        shape_text(n, lab, size=11.5, color=(WHITE if is_os else NAVY), bold=True)
    text(s, cx - 1.6, cy - R - 0.38, 3.2, 0.3, "THE STARTUP ROOF ECOSYSTEM", size=8.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER, spacing=1.2)

    # definitions
    dx = 6.85
    rect(s, dx, 2.2, 0.06, 1.55, fill=NAVY)
    text(s, dx + 0.3, 2.18, 5.6, 0.35, [[B("STARTUP ROOF", color=NAVY, spacing=1.2), ("  =  the wider ecosystem", {"color": MUTED})]], size=13)
    text(s, dx + 0.3, 2.6, 5.6, 1.2, "The community, programmes, relationships and services that surround a founder all year — online and in person.", size=13, color=TEXT, line=1.25)
    rect(s, dx, 3.98, 0.06, 1.55, fill=BLUE)
    text(s, dx + 0.3, 3.96, 5.6, 0.35, [[B("FOUNDER OS", color=BLUE, spacing=1.2), ("  =  the digital workspace", {"color": MUTED})]], size=13)
    text(s, dx + 0.3, 4.38, 5.6, 1.2, "The guided online experience inside the ecosystem: a founder’s profile, goals, next steps, resources and connections in one place.", size=13, color=TEXT, line=1.25)
    text(s, dx, 5.85, 5.9, 0.9,
         [[B("In short: ", color=NAVY), ("an ecosystem around the founder, and a digital workspace for the journey.", {"color": TEXT})]],
         size=13.5, line=1.25)
    footer(s, 3)
    notes(s, "This is the most important distinction in the deck. Startup Roof is the ecosystem — people, programmes, services and relationships. "
             "Founder OS is the software layer that helps a founder use that ecosystem: it organises their goals and next steps and routes them to the right resources and people. "
             "Neither replaces the other: software makes the ecosystem easier to navigate, and the ecosystem gives the software something real to connect to.")


def s_who():
    global s
    s = blank(prs)
    header(s, 4, "Who it is for", "Built for founders. Strengthened by the people around them.")
    p = rect(s, ML, 2.15, 4.6, 4.6, fill=NAVY)
    text(s, ML + 0.35, 2.4, 4, 0.3, "PRIMARY  ·  FOUNDERS", size=9.5, color=MID, bold=True, spacing=1.3)
    prim = [("Aspiring entrepreneurs", "Want to test an idea and understand where to start."),
            ("Early-stage founders", "Need validation, structure and their first customers."),
            ("Founders seeking growth", "Need operational support, specialists and sharper go-to-market.")]
    for i, (a, b) in enumerate(prim):
        y = 2.8 + i * 1.12
        text(s, ML + 0.35, y, 3.9, 0.35, a, size=15, color=WHITE, bold=True)
        text(s, ML + 0.35, y + 0.4, 3.9, 0.7, b, size=12, color="C9D6E8", line=1.2)
    text(s, ML + 0.35, 6.15, 4, 0.5, "Technology startups and commercially viable businesses in consumer, wellness, retail and services.", size=9.5, color=MID, line=1.2)

    text(s, 5.65, 2.2, 6, 0.3, "SUPPORTING ECOSYSTEM  ·  WHY THEY TAKE PART", size=9.5, color=BLUE, bold=True, spacing=1.3)
    sup = [("Mentors and specialists", "Share experience in a structured way and meet founders who fit their expertise."),
           ("Investors and angels", "Engage with better-prepared founders and follow sectors of interest."),
           ("Universities and incubators", "Extend support to student and alumni founders beyond a single programme."),
           ("Corporate innovation teams", "Meet early ventures and run practical challenges or pilots."),
           ("Technology and service providers", "Offer relevant tools and services to founders who need them.")]
    for i, (a, b) in enumerate(sup):
        y = 2.62 + i * 0.83
        line(s, 5.65, y, W - ML, y, color=LINE)
        text(s, 5.65, y + 0.14, 2.75, 0.6, a, size=13.5, color=NAVY, bold=True, line=1.1)
        text(s, 8.55, y + 0.14, 4.1, 0.65, b, size=12, color=TEXT, line=1.2)
    footer(s, 4)
    notes(s, "Founders are the reason the ecosystem exists; everyone else participates because they get something useful from being near the right founders. "
             "Stress the breadth: this is not only for venture-backed tech. Consumer brands, wellness, retail and service businesses are equally in scope, and many will never need to raise venture funding.")


def s_journey():
    global s
    s = blank(prs)
    header(s, 5, "The founder journey", "Every founder enters at a different point.")
    stages = [
        ("Explore", "Is this idea worth pursuing?", "Framing, problem discovery, peer conversation", "Write down the problem and who has it"),
        ("Validate", "Will people actually pay for this?", "Interview guides, clinics, mentor feedback", "Speak to 10–20 potential customers"),
        ("Build", "What is the simplest version to offer?", "Product, setup, legal and tool guidance", "Define a first offer and test it"),
        ("Launch", "How do I reach my first customers?", "Go-to-market, pricing, brand and sales support", "Run one focused acquisition test"),
        ("Grow", "How do I grow without breaking things?", "Operations, specialists, systems and hiring", "Fix the biggest bottleneck"),
        ("Capital", "Should I raise — and am I ready?", "Readiness review, pitch feedback, introductions", "Decide if external funding fits"),
    ]
    lx, top = ML + 1.45, 2.15
    cw = (W - ML - lx) / 6
    labels = ["Founder question", "Support needed", "Useful next action"]
    for i, (st, q, sup_, act) in enumerate(stages):
        x = lx + i * cw
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        ch = rect(s, x, top, cw + 0.08, 0.62, fill=(BLUE if i < 5 else NAVY), shape=shp)
        shape_text(ch, st if i < 5 else "Prepare for capital*", size=12 if i < 5 else 10.5, color=WHITE, bold=True)
        for r, val in enumerate([q, sup_, act]):
            y = top + 0.85 + r * 1.12
            text(s, x + 0.12, y, cw - 0.2, 1.1, val, size=11.5 if r == 0 else 11, color=(NAVY if r == 0 else TEXT),
                 bold=(r == 0), line=1.18)
    for r, lab in enumerate(labels):
        y = top + 0.85 + r * 1.12
        line(s, ML, y - 0.12, W - ML, y - 0.12, color=LINE)
        text(s, ML, y, 1.35, 0.8, lab, size=9.5, color=BLUE, bold=True, spacing=0.6, line=1.15)
    rect(s, ML, 6.2, W - 2 * ML, 0.6, fill=PALE)
    text(s, ML + 0.2, 6.25, W - 2 * ML - 0.4, 0.5,
         [[B("Paths differ. ", color=NAVY), ("Founders join at any stage and may loop back. *Many good businesses grow from revenue and never need external capital — this stage applies only where appropriate.", {})]],
         size=11, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 5)
    notes(s, "Walk one column top to bottom: the question a founder is really asking, the support that helps, and one concrete next action. "
             "The point is that Startup Roof meets founders where they are. A founder with an established consumer brand may start at Grow; a student may start at Explore. "
             "Capital is the last stage and explicitly optional.")


def s_layers():
    global s
    s = blank(prs)
    header(s, 6, "The proposed model", "Four layers, from open access to selective support.", tag="PROPOSED MODEL")
    layers = [
        ("4", "Investment connections and future pathways", "Investor-readiness support and relevant introductions. Any future investment activity would need its own structure and approvals.", NAVY, 2.2),
        ("3", "Selective incubation", "More intensive support for suitable founders, based on fit, delivery capacity and separately agreed commercial terms.", "17407F", 3.1),
        ("2", "Paid tools and support", "Deeper tools, potential AI-assisted guidance, specialist sessions and structured programmes.", "1A55AC", 4.0),
        ("1", "Free founder access", "Community, selected resources, discovery and introductory activities — the open front door.", BLUE, 4.9),
    ]
    cx0 = ML + 2.6
    for i, (n, t, d, col, wdt) in enumerate(layers):
        y = 2.2 + i * 1.05
        rr = rect(s, cx0 - wdt / 2, y, wdt, 0.88, fill=col)
        shape_text(rr, [[B(n, size=20)]], size=20, color=WHITE)
        text(s, 5.95, y + 0.05, 6.7, 0.35, t, size=15, color=NAVY, bold=True)
        text(s, 5.95, y + 0.42, 6.7, 0.55, d, size=11.5, color=TEXT, line=1.18)
        line(s, cx0 + wdt / 2 + 0.1, y + 0.44, 5.8, y + 0.44, color=LINE, w=0.75, dash=True)
    text(s, ML, 6.48, 11.9, 0.4,
         [[B("Fewer founders, deeper support, as you move up. ", color=NAVY),
           ("Pricing, terms and eligibility will be shaped during the pilot — no fees or equity terms are fixed today.", {"color": MUTED})]],
         size=11.5)
    footer(s, 6)
    notes(s, "Read the layers bottom-up. The free layer creates reach and trust. The paid layer is where useful tools and expert time are offered on fair terms. "
             "Incubation is selective because intensive support is expensive to deliver well. The top layer is about investor readiness and introductions — "
             "Startup Roof is not operating a fund, and anything like that would require its own structure, approvals and terms.")


def s_founder_os():
    global s
    s = blank(prs)
    header(s, 7, "Founder OS", "A digital workspace for the founder journey.", tag="PLANNED CAPABILITY")
    # illustrative UI concept
    ux, uy, uw, uh = ML, 2.15, 6.15, 4.3
    rect(s, ux, uy, uw, uh, fill=WHITE, line="C9D6E8", lw=1)
    rect(s, ux, uy, uw, 0.32, fill=PALE)
    for k in range(3):
        rect(s, ux + 0.14 + k * 0.17, uy + 0.11, 0.1, 0.1, fill="B8CFEE", shape=MSO_SHAPE.OVAL)
    text(s, ux + 0.8, uy + 0.07, 3, 0.2, "Founder OS", size=8.5, color=MUTED, bold=True)
    rect(s, ux, uy + 0.32, 1.35, uh - 0.32, fill=NAVY)
    for k, lab in enumerate(["Dashboard", "My startup", "Next steps", "Learn", "People", "Opportunities"]):
        text(s, ux + 0.18, uy + 0.55 + k * 0.38, 1.15, 0.25, lab, size=9, color=(WHITE if k == 0 else "9FB2CC"), bold=(k == 0))
    ix = ux + 1.6
    text(s, ix, uy + 0.5, 4.3, 0.3, "Your current stage", size=8.5, color=MUTED)
    text(s, ix, uy + 0.72, 4.3, 0.35, "Validate", size=16, color=NAVY, bold=True)
    for k in range(6):
        rect(s, ix + k * 0.7, uy + 1.15, 0.62, 0.08, fill=(BLUE if k <= 1 else "D9E5F6"))
    rect(s, ix, uy + 1.45, 4.35, 1.2, fill=PALE)
    text(s, ix + 0.15, uy + 1.55, 4, 0.25, "SUGGESTED NEXT STEP", size=7.5, color=BLUE, bold=True, spacing=0.8)
    text(s, ix + 0.15, uy + 1.8, 4.1, 0.5, "Run five customer conversations using the interview guide", size=11, color=NAVY, bold=True, line=1.1)
    bt = rect(s, ix + 0.15, uy + 2.25, 1.25, 0.28, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    shape_text(bt, "Open guide", size=8, color=WHITE, bold=True)
    for k, (a, b) in enumerate([("Resource", "Customer interview template"), ("Session", "Validation clinic — book a slot"), ("Connect", "Request a mentor in your sector")]):
        y = uy + 2.85 + k * 0.42
        text(s, ix, y, 1.0, 0.3, a.upper(), size=7.5, color=BLUE, bold=True, spacing=0.6)
        text(s, ix + 1.0, y - 0.02, 3.3, 0.3, b, size=9.5, color=TEXT)
        line(s, ix, y + 0.32, ix + 4.35, y + 0.32, color="E3ECF8")
    text(s, ux, uy + uh + 0.1, uw, 0.25, "Illustrative product concept — not a live screen.", size=9, color=MUTED, italic=True)

    groups = [
        ("My startup", "Founder profile · stage and goals · personal dashboard"),
        ("Next steps", "Tasks · resources and templates · learning progress"),
        ("People and support", "Mentor and specialist discovery · connection requests · service requests"),
        ("Opportunities and progress", "Relevant opportunities · activity and progress tracking"),
    ]
    gx = 7.25
    for i, (g, d) in enumerate(groups):
        y = 2.15 + i * 1.08
        rect(s, gx, y + 0.05, 0.4, 0.4, fill=(BLUE if i % 2 == 0 else NAVY), shape=MSO_SHAPE.OVAL)
        text(s, gx, y + 0.1, 0.4, 0.3, str(i + 1), size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(s, gx + 0.6, y, 4.8, 0.35, g, size=15, color=NAVY, bold=True)
        text(s, gx + 0.6, y + 0.38, 4.8, 0.6, d, size=11.5, color=TEXT, line=1.2)
    text(s, gx, 6.5, 5.4, 0.4, "Capabilities will be prioritised with pilot founders and released step by step.", size=11, color=MUTED, line=1.2)
    footer(s, 7)
    notes(s, "Founder OS is planned, not live. The screen on the left is an illustrative concept to make the idea concrete. "
             "Its job is simple: know where a founder is, suggest a sensible next step, and route them to the right resource, session or person. "
             "The four capability groups will be prioritised with pilot founders — we will build the smallest useful version first.")


def s_scenario():
    global s
    s = blank(prs)
    header(s, 8, "The founder experience", "A founder joins with an idea and needs to validate demand.", tag="ILLUSTRATIVE SCENARIO", title_size=26)
    card = rect(s, ML, 2.15, 3.35, 4.65, fill=PALE)
    text(s, ML + 0.3, 2.4, 2.8, 0.3, "THE SCENARIO", size=9.5, color=BLUE, bold=True, spacing=1.2)
    text(s, ML + 0.3, 2.8, 2.85, 2.2,
         "A first-time founder in Dubai wants to launch a healthy-snack brand. The recipe is ready and the enthusiasm is real, but there is no evidence yet that customers will buy.",
         size=13.5, color=NAVY, line=1.3)
    text(s, ML + 0.3, 5.05, 2.85, 1.5,
         "This is an example of how the experience is designed to work — not a real customer case.",
         size=11, color=MUTED, line=1.25, italic=True)
    steps = [
        ("Complete a profile", "Stage, sector, goals and time available."),
        ("Identify the challenge", "“I don’t know if anyone will pay for this.”"),
        ("Get a relevant next step", "An interview guide and a short validation task."),
        ("Attend a clinic or ask for help", "A group validation clinic or a mentor request."),
        ("Record progress", "Notes from 12 conversations; pricing feedback logged."),
        ("Return for the next step", "Move on to testing a first offer."),
    ]
    sx, sy = 4.4, 2.2
    colw = 2.85
    for i, (a, b) in enumerate(steps):
        col, row = i % 3, i // 3
        x = sx + col * colw
        y = sy + row * 2.3
        c = rect(s, x, y, 0.55, 0.55, fill=(BLUE if i < 5 else NAVY), shape=MSO_SHAPE.OVAL)
        shape_text(c, str(i + 1), size=15, color=WHITE, bold=True)
        if col < 2:
            line(s, x + 0.65, y + 0.27, x + colw - 0.1, y + 0.27, color="B8CFEE", w=1.25, arrow=True)
        text(s, x, y + 0.72, colw - 0.3, 0.6, a, size=13.5, color=NAVY, bold=True, line=1.1)
        text(s, x, y + 1.3, colw - 0.35, 0.8, b, size=11.5, color=TEXT, line=1.2)
    footer(s, 8)
    notes(s, "Make clear this is a scenario, not a customer story. It shows how the ecosystem and the software work together: "
             "Founder OS captures the profile and suggests the next step; Startup Roof supplies the clinic, the mentor and the community. "
             "The founder records progress and comes back — the loop is the product. Success is measured by whether founders take and complete useful next steps.")


def s_ai_human():
    global s
    s = blank(prs)
    header(s, 9, "AI guidance with human support", "Digital guidance, strengthened by human experience.", tag="PROPOSED APPROACH")
    colx = [ML, 8.08]
    cols = [("AI-ASSISTED SUPPORT", "Helps founders prepare and keep momentum",
             ["Organising information", "Exploring questions", "Drafting plans", "Suggesting resources", "Tracking next steps"], BLUE),
            ("HUMAN SUPPORT", "Brings judgement and real-world context",
             ["Mentoring", "Specialist judgement", "Practical feedback", "Complex business decisions", "Sector expertise"], NAVY)]
    for (title, sub, its, col), x in zip(cols, colx):
        rect(s, x, 2.2, 4.6, 0.08, fill=col)
        text(s, x, 2.45, 4.6, 0.3, title, size=10, color=col, bold=True, spacing=1.3)
        text(s, x, 2.78, 4.6, 0.4, sub, size=14, color=NAVY, bold=True)
        for k, it in enumerate(its):
            y = 3.4 + k * 0.6
            line(s, x, y, x + 4.6, y, color=LINE)
            text(s, x, y + 0.13, 4.6, 0.4, it, size=14, color=TEXT)
    mc = rect(s, 5.55, 2.85, 2.25, 2.25, fill=PALE, shape=MSO_SHAPE.OVAL)
    shape_text(mc, [[B("The founder\ndecides", color=NAVY)]], size=15, color=NAVY)
    text(s, ML, 6.5, 12, 0.35,
         [[B("Principle: ", color=NAVY), ("AI prepares, people judge, founders decide. AI features will be tested with pilot founders before release; none is presented here as live.", {"color": MUTED})]],
         size=11.5)
    footer(s, 9)
    notes(s, "AI is useful for structure and speed: organising notes, drafting a plan, suggesting what to read next and reminding founders what they committed to. "
             "It is not a substitute for a mentor who has run a business in that sector. The design principle is that AI prepares, people judge and the founder decides. "
             "No AI integration is live today; features will be tested for usefulness and accuracy in the pilot.")


def s_community():
    global s
    s = blank(prs)
    header(s, 10, "Community and meaningful connections", "Relevant connections, measured by what happens next.", title_size=27)
    flow = ["Request", "Consent", "Introduction", "Conversation", "Follow-up", "Outcome noted"]
    fx, fy = ML, 2.35
    fw = (W - 2 * ML) / 6
    for i, f in enumerate(flow):
        x = fx + i * fw
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        c = rect(s, x, fy, fw + 0.06, 0.6, fill=(PALE if i not in (1, 5) else (BLUE if i == 1 else NAVY)), shape=shp)
        shape_text(c, f, size=12, color=(NAVY if i not in (1, 5) else WHITE), bold=True)
    text(s, ML, 3.08, 12, 0.3, "How a useful connection is designed to work — each introduction is opt-in on both sides and followed up.", size=11, color=MUTED)

    text(s, ML, 3.7, 6, 0.3, "WAYS PEOPLE MAY CONNECT", size=9.5, color=BLUE, bold=True, spacing=1.3)
    mech = [("Founder introductions", "matched by need, stage or sector"), ("Peer groups", "small cohorts facing similar problems"),
            ("Office hours", "short, bookable time with experts"), ("Mentor sessions", "focused conversations with follow-up"),
            ("Small gatherings", "in-person meetups in the UAE"), ("Community discussions", "questions answered by peers"),
            ("Accountability", "check-ins on commitments made")]
    for i, (a, b) in enumerate(mech):
        col, row = i % 2, i // 2
        x = ML + col * 3.65
        y = 4.08 + row * 0.62
        text(s, x, y, 3.55, 0.6, [[B(a, color=NAVY)], [(b, {"color": MUTED, "size": 11})]], size=12.5, line=1.1)
    pp = rect(s, 8.3, 3.7, 4.38, 3.05, fill=NAVY)
    text(s, 8.6, 3.92, 3.8, 0.3, "PRODUCT PRINCIPLES", size=9.5, color=MID, bold=True, spacing=1.3)
    pr = [("Consent first", "No introduction without agreement from both sides."),
          ("Founder-controlled visibility", "Founders choose what their profile shows, and to whom."),
          ("Quality over size", "We track completed, useful interactions — not directory size.")]
    for k, (a, b) in enumerate(pr):
        y = 4.32 + k * 0.8
        text(s, 8.6, y, 3.85, 0.8, [[B(a, color=WHITE)], [(b, {"color": "C9D6E8", "size": 11})]], size=13, line=1.15)
    footer(s, 10)
    notes(s, "Most networks measure how many people are listed. We want to measure whether introductions led to a real conversation and a useful outcome. "
             "Consent and profile visibility are product principles from day one: founders decide what is shared, and both sides opt in before any introduction.")


def s_academy():
    global s
    s = blank(prs)
    header(s, 11, "Founder Academy", "Practical learning that ends in action.", tag="PROPOSED THEMES")
    themes = ["Idea validation", "Business models", "AI and automation", "Go-to-market and customer acquisition",
              "Pricing and unit economics", "Pitch readiness", "Legal and operational foundations", "Consumer-brand and retail development"]
    for i, t in enumerate(themes):
        col, row = i % 2, i // 2
        x = ML + col * 3.55
        y = 2.25 + row * 0.98
        line(s, x, y, x + 3.3, y, color=NAVY if row == 0 else LINE, w=1.0)
        text(s, x, y + 0.14, 0.5, 0.3, f"{i + 1:02d}", size=11, color=BLUE, bold=True)
        text(s, x + 0.5, y + 0.12, 2.8, 0.75, t, size=13.5, color=NAVY, bold=True, line=1.1)
    text(s, ML, 6.25, 6.9, 0.6, "Themes are a starting point; sessions and depth will be shaped with founders and contributors during the pilot.",
         size=11, color=MUTED, line=1.25)
    # cycle
    ccx, ccy = 10.25, 4.25
    ring = rect(s, ccx - 1.75, ccy - 1.75, 3.5, 3.5, fill=None, line="B8CFEE", lw=2, shape=MSO_SHAPE.OVAL)
    cyc = [("Learn", ccx, ccy - 1.75), ("Apply", ccx + 1.75, ccy), ("Feedback", ccx, ccy + 1.75), ("Next action", ccx - 1.75, ccy)]
    for k, (lab, nx, ny) in enumerate(cyc):
        n = rect(s, nx - 0.65, ny - 0.27, 1.3, 0.54, fill=(BLUE if k != 3 else NAVY), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        shape_text(n, lab, size=11.5, color=WHITE, bold=True)
    text(s, ccx - 0.95, ccy - 0.5, 1.9, 1.0, [[B("Every session", color=NAVY)], [("ends with a task the founder applies to their own business", {"color": TEXT, "size": 10.5})]],
         size=12, align=PP_ALIGN.CENTER, line=1.15)
    footer(s, 11)
    notes(s, "The Academy is not a video library. Each theme follows a simple loop: learn a concept, apply it to your own business, get feedback from a mentor or peers, and leave with a next action. "
             "These eight themes are proposed — no course catalogue exists yet. Consumer-brand and retail development is included deliberately because many founders we want to serve are not building software.")


def s_services():
    global s
    s = blank(prs)
    header(s, 12, "Expert services and partner tools", "The right specialist or tool, at the right moment.", title_size=27)
    cats = ["Legal", "Accounting and finance", "Branding and design", "Technology and automation",
            "Marketing and sales", "Customer communication", "Business setup", "Product and retail guidance"]
    text(s, ML, 2.2, 5, 0.3, "POTENTIAL SERVICE CATEGORIES", size=9.5, color=BLUE, bold=True, spacing=1.3)
    for i, c in enumerate(cats):
        col, row = i % 2, i // 2
        x = ML + col * 2.6
        y = 2.6 + row * 0.82
        b = rect(s, x, y, 2.45, 0.68, fill=PALE)
        shape_text(b, c, size=12, color=NAVY, bold=True, align=PP_ALIGN.LEFT, margin=0.18)
    routes = [("Expert sessions", "Short clinics and one-to-one sessions with experienced practitioners."),
              ("Partner-provided services", "Defined packages from vetted providers, with clear scope before work starts."),
              ("Integrated or white-label tools", "Useful software inside Founder OS, where commercial rights allow.")]
    rx = 6.3
    text(s, rx, 2.2, 6, 0.3, "THREE POSSIBLE DELIVERY ROUTES", size=9.5, color=BLUE, bold=True, spacing=1.3)
    for k, (a, b) in enumerate(routes):
        y = 2.6 + k * 1.1
        rect(s, rx, y, 0.06, 0.9, fill=[BLUE, "1A55AC", NAVY][k])
        text(s, rx + 0.25, y + 0.02, 6, 0.35, a, size=15, color=NAVY, bold=True)
        text(s, rx + 0.25, y + 0.42, 6.1, 0.5, b, size=12, color=TEXT, line=1.2)
    rect(s, rx, 5.98, W - ML - rx, 0.8, fill=NAVY)
    text(s, rx + 0.25, 6.08, W - ML - rx - 0.5, 0.65,
         "Packages, providers, integrations and fulfilment will be validated during the pilot. No supplier is presented here as a confirmed partner.",
         size=11.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, line=1.2)
    footer(s, 12)
    notes(s, "Founders repeatedly need the same specialist help — legal, accounting, setup, branding, marketing. We want to make that help easier to find and fairer to buy. "
             "Three delivery routes are being explored. Which providers, at what price and with what quality controls, will be decided during the pilot. No partners are confirmed.")


def s_insights():
    global s
    s = blank(prs)
    header(s, 13, "Startup insights, media and opportunities", "What matters to a founder this week — with context.", title_size=27)
    rcx, rcy = 3.55, 4.5
    for k, rr in enumerate([2.1, 1.5, 0.9]):
        rect(s, rcx - rr, rcy - rr, 2 * rr, 2 * rr, fill=[PALE, "DCE8FB", "C4D8F6"][k], shape=MSO_SHAPE.OVAL)
    line(s, rcx - 2.1, rcy, rcx + 2.1, rcy, color="B8CFEE", w=0.75)
    line(s, rcx, rcy - 2.1, rcx, rcy + 2.1, color="B8CFEE", w=0.75)
    cc = rect(s, rcx - 0.48, rcy - 0.48, 0.96, 0.96, fill=NAVY, shape=MSO_SHAPE.OVAL)
    shape_text(cc, "Radar", size=11.5, color=WHITE, bold=True)
    quad = [("Startup updates", rcx - 1.85, rcy - 1.55), ("Relevant events", rcx + 0.25, rcy - 1.55),
            ("Funding and programme opportunities", rcx - 1.85, rcy + 0.75), ("Founder interviews and media", rcx + 0.25, rcy + 0.75)]
    for lab, qx, qy in quad:
        text(s, qx, qy, 1.6, 0.8, lab, size=11, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.1)
    for dx_, dy_ in [(0, -1.2), (1.2, 0), (0, 1.2), (-1.2, 0)]:
        rect(s, rcx + dx_ - 0.07, rcy + dy_ - 0.07, 0.14, 0.14, fill=BLUE, shape=MSO_SHAPE.OVAL)
    text(s, 6.4, 2.2, 6.3, 0.9,
         [[B("Startup Roof Radar", color=NAVY)], [("A planned, curated view of updates, events and opportunities — starting with the UAE and widening over time.", {"color": TEXT, "size": 12.5})]],
         size=16, line=1.2)
    prin = [("Relevance", "Filtered by a founder’s stage, sector and location."), ("Attribution", "Every item credits and links to its original source."),
            ("Freshness", "Dated, and removed when no longer current."), ("Founder context", "A short note on why it matters and what to do next.")]
    for k, (a, b) in enumerate(prin):
        y = 3.35 + k * 0.72
        line(s, 6.4, y, W - ML, y, color=LINE)
        text(s, 6.4, y + 0.14, 2.1, 0.4, a, size=13.5, color=BLUE, bold=True)
        text(s, 8.5, y + 0.14, 4.2, 0.5, b, size=12, color=TEXT, line=1.15)
    text(s, 6.4, 6.3, 6.3, 0.5, "Coverage will be selective and grow step by step; it is not intended to be a complete news service.", size=11, color=MUTED, line=1.2)
    footer(s, 13)
    notes(s, "Radar is the working name for a curated feed of what is useful to founders now: news, events, programmes and funding calls. "
             "Our differentiator is not volume but relevance and context — each item credited to its source, dated, and annotated with why it matters. "
             "Founder interviews and media help share practical lessons. Coverage starts UAE-first and will be selective, not comprehensive.")


def s_incubation():
    global s
    s = blank(prs)
    header(s, 14, "Incubation and commercial readiness", "Selective, structured support where it can make a difference.", title_size=26)
    steps = ["Understand the founder’s needs", "Assess suitability", "Define a practical support plan",
             "Match mentors and specialists", "Agree milestones", "Review progress"]
    sx, sy = ML, 2.3
    sw = (W - 2 * ML) / 6
    for i, st in enumerate(steps):
        x = sx + i * sw
        y = sy + (5 - i) * 0.0
        num = rect(s, x, sy, 0.5, 0.5, fill=(NAVY if i == 1 else BLUE), shape=MSO_SHAPE.OVAL)
        shape_text(num, str(i + 1), size=13, color=WHITE, bold=True)
        if i < 5:
            line(s, x + 0.6, sy + 0.25, x + sw - 0.08, sy + 0.25, color="B8CFEE", w=1.25, arrow=True)
        text(s, x, sy + 0.65, sw - 0.2, 0.75, st, size=12.5, color=NAVY, bold=True, line=1.12)
    gate = rect(s, sx + sw - 0.05, sy + 1.45, sw - 0.1, 0.3, fill=PALE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    shape_text(gate, "Selection point", size=9, color=BLUE, bold=True)

    text(s, ML, 4.25, 8, 0.3, "EXAMPLE SUPPORT PLANS", size=9.5, color=BLUE, bold=True, spacing=1.3)
    data = [["", "Technology startup", "Consumer or service business"],
            ["Typical focus", "Product–market fit, early traction, technical choices", "Unit economics, channels, supply and operations"],
            ["Example milestone", "First ten paying customers on a defined plan", "Margin-positive pilot with one retail or online channel"],
            ["Typical specialists", "Product, AI/automation, go-to-market", "Brand, retail, sourcing, finance"]]
    table(s, ML, 4.6, 8.3, 1.9, data, [2.0, 3.15, 3.15], size=11, header_size=10.5, row_h=0.45)
    rect(s, 9.3, 4.25, W - ML - 9.3, 2.5, fill=NAVY)
    text(s, 9.55, 4.45, 3.2, 2.2,
         [[B("Intensive support is selective.", color=WHITE)],
          {"runs": [("Places depend on fit, the founder’s commitment and our delivery capacity. Terms are agreed separately with each business.", {"color": "C9D6E8", "size": 11.5})], "before": 6}],
         size=14, line=1.25)
    footer(s, 14)
    notes(s, "This is the deeper end of the model. The process starts by understanding the founder, then an honest suitability check — the selection point — before any plan is agreed. "
             "Plans have milestones and regular reviews. The examples show that the same process works for a tech startup and a consumer or service business; the specialists and milestones differ. "
             "Intensive support is limited by delivery capacity, so it will be selective.")


def s_investors():
    global s
    s = blank(prs)
    header(s, 15, "Value for investors", "Better-prepared founders. More useful conversations.")
    vals = [("Better founder preparation", "Founders work through readiness before an introduction."),
            ("More structured information", "Consistent profiles, stage and progress — shared with consent."),
            ("Relevant discovery", "Filter by sector, stage and geography."),
            ("Sector-specific engagement", "Join sessions in areas you know well."),
            ("More useful introductions", "Fewer, better-matched conversations."),
            ("Ongoing ecosystem relationships", "Stay close to founders and mentors over time.")]
    for i, (a, b) in enumerate(vals):
        y = 2.2 + i * 0.75
        rect(s, ML, y + 0.08, 0.12, 0.12, fill=BLUE)
        text(s, ML + 0.3, y, 5.3, 0.35, a, size=13.5, color=NAVY, bold=True)
        text(s, ML + 0.3, y + 0.33, 5.3, 0.35, b, size=11.5, color=MUTED)
    data = [["STARTUP ROOF AIMS TO", "STARTUP ROOF DOES NOT"],
            ["Help founders become investor-ready", "Operate its own investment fund today"],
            ["Make relevant, consented introductions", "Promise deal flow, exclusivity or returns"],
            ["Share structured founder information", "Recommend or endorse specific investments"],
            ["Host sector sessions with investors", "Assume every founder should raise capital"]]
    table(s, 6.55, 2.2, W - ML - 6.55, 3.6, data, [3.07, 3.06], size=12, header_size=10.5, first_col_bold=False, row_h=0.7)
    text(s, 6.55, 6.05, W - ML - 6.55, 0.7,
         "Any future investment activity would require its own structure, approvals and terms.", size=11, color=MUTED, line=1.2)
    footer(s, 15)
    notes(s, "For investors the value is quality, not volume: founders who have done the groundwork and introductions that match their thesis. "
             "Be explicit about what we are not: Startup Roof is not operating a fund, does not promise returns or exclusive deals, and does not endorse investments. "
             "Investor engagement today means sessions, mentoring and relevant introductions.")


def s_partners():
    global s
    s = blank(prs, PALE)
    header(s, 16, "Value for ecosystem partners", "Collaborations that are useful to founders and to partners.", title_size=27)
    why = ["Reach relevant founders", "Deliver education or specialist support", "Test useful services",
           "Build long-term relationships", "Support entrepreneurship programmes", "Explore commercial distribution"]
    text(s, ML, 2.2, 4, 0.3, "WHY PARTNERS MAY TAKE PART", size=9.5, color=BLUE, bold=True, spacing=1.3)
    for i, w_ in enumerate(why):
        y = 2.62 + i * 0.63
        text(s, ML, y, 0.45, 0.4, f"{i + 1:02d}", size=12, color=BLUE, bold=True)
        text(s, ML + 0.5, y, 3.6, 0.5, w_, size=13.5, color=NAVY, bold=True)
    data = [["COLLABORATION FORMAT", "SUITS", "WHAT IT COULD LOOK LIKE"],
            ["Workshops", "Specialists, firms", "A practical 90-minute session on a founder topic"],
            ["Office hours", "Mentors, advisers", "Regular bookable slots with short follow-up"],
            ["University programmes", "Universities, incubators", "Support for student and alumni founders"],
            ["Tool or service pilots", "Technology and service providers", "A limited trial with feedback from founders"],
            ["Community activities", "Corporates, venues", "Hosting a meetup or a themed founder evening"],
            ["Content collaboration", "Media, experts", "Guides, interviews or co-created resources"]]
    table(s, 4.85, 2.2, W - ML - 4.85, 4.3, data, [2.25, 2.35, 3.23], size=11.5, header_size=10, row_fills=(WHITE,), row_h=0.6)
    text(s, 4.85, 6.6, 7.8, 0.3, "These are possible formats to explore together; no partnerships are presented here as signed.", size=11, color=MUTED)
    footer(s, 16)
    notes(s, "Partners include universities, incubators, corporates, professional firms and technology providers. "
             "The table offers concrete, low-commitment ways to start — a workshop, office hours or a small pilot — so both sides can learn whether the collaboration is useful. "
             "No partnerships are signed; this slide is an invitation to explore.")


def s_commercial():
    global s
    s = blank(prs)
    header(s, 17, "A sustainable commercial model", "Free access can create reach.\nUseful services must create sustainable value.", title_size=27)
    flow = [("Free access", "Community, resources and discovery build reach and trust.", PALE, NAVY),
            ("Optional paid value", "Founders and partners pay only for what is clearly useful.", BLUE, WHITE),
            ("Delivery costs", "Expert time, platform, events and support all have real costs.", "DCE8FB", NAVY),
            ("Unit economics", "To validate: each service must cover its costs and contribute to the whole.", NAVY, WHITE)]
    fw = (W - 2 * ML - 0.6) / 4
    for i, (a, b, fc, tc) in enumerate(flow):
        x = ML + i * (fw + 0.2)
        r_ = rect(s, x, 2.35, fw, 1.65, fill=fc)
        text(s, x + 0.2, 2.5, fw - 0.4, 0.35, a, size=13.5, color=tc, bold=True)
        text(s, x + 0.2, 2.95, fw - 0.4, 1.0, b, size=11.5, color=(TEXT if tc == NAVY else ("DCE6F5" if fc == NAVY else WHITE)), line=1.2)
    text(s, ML, 4.35, 6, 0.3, "REVENUE HYPOTHESES TO TEST", size=9.5, color=BLUE, bold=True, spacing=1.3)
    rev = ["Paid memberships or tools", "Specialist support packages", "Structured programmes",
           "Partner and service revenue arrangements", "Corporate or university programmes",
           "Sponsorship of suitable activities", "Selective incubation arrangements, where separately agreed"]
    for i, r_ in enumerate(rev):
        col, row = (0, i) if i < 4 else (1, i - 4)
        x = ML + col * 6.1
        y = 4.72 + row * 0.5
        rect(s, x, y + 0.1, 0.1, 0.1, fill=BLUE)
        text(s, x + 0.25, y, 5.7, 0.45, r_, size=12.5, color=TEXT)
    text(s, ML + 6.1, 6.25, 6, 0.6, "No financial forecast is presented. Pricing and margins will be tested in the pilot before any are published.",
         size=11, color=MUTED, line=1.2)
    footer(s, 17)
    notes(s, "The logic is simple. Free access is how founders discover us — but free has a cost and must be funded. "
             "Revenue should come from things people choose to pay for because they are useful: tools, specialist packages, programmes and partner arrangements. "
             "Every one of these is a hypothesis. The pilot will measure delivery costs and willingness to pay before we commit to a model.")


def s_roadmap():
    global s
    s = blank(prs)
    header(s, 18, "Trial, validation and roadmap", "Shape, test, then expand what works.")
    phases = [("NOW", "Shape and test", ["Founder interviews", "Participant contributions", "Service selection", "Platform requirements"], BLUE),
              ("NEXT", "Deliver a small pilot", ["Useful sessions", "Onboarding and follow-up", "Growth experiments", "Measure delivery costs"], "1A55AC"),
              ("THEN", "Expand what works", ["Repeat strong services", "Improve the platform", "Formalise responsibilities", "Build partnerships"], NAVY)]
    pw = (W - 2 * ML - 0.4) / 3
    for i, (tag, ttl, its, col) in enumerate(phases):
        x = ML + i * (pw + 0.2)
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        c = rect(s, x, 2.2, pw + (0.15 if i < 2 else 0), 0.55, fill=col, shape=shp)
        shape_text(c, [[B(tag, spacing=1.5), ("   " + ttl, {})]], size=13, color=WHITE, align=PP_ALIGN.LEFT, margin=0.3)
        for k, it in enumerate(its):
            y = 2.98 + k * 0.42
            rect(s, x + 0.05, y + 0.1, 0.09, 0.09, fill=col)
            text(s, x + 0.28, y, pw - 0.4, 0.35, it, size=12.5, color=TEXT)
    rect(s, ML, 4.8, W - 2 * ML, 2.0, fill=PALE)
    text(s, ML + 0.3, 5.0, 4, 0.3, "MILESTONES WE ARE WORKING TOWARDS", size=9.5, color=BLUE, bold=True, spacing=1.3)
    ms = [("10,000", "free users", "Threshold for a public launch event"),
          ("March 2027", "tentative", "Target timing, subject to readiness"),
          ("100,000", "free users", "Wider, longer-term ambition")]
    mw = (W - 2 * ML - 0.6) / 3
    for i, (a, b, c_) in enumerate(ms):
        x = ML + 0.3 + i * mw
        text(s, x, 5.38, mw - 0.2, 0.55, [[B(a, color=NAVY, size=24), ("  " + b, {"color": MUTED, "size": 12})]], size=24)
        text(s, x, 5.95, mw - 0.3, 0.4, c_, size=12, color=TEXT)
    text(s, ML + 0.3, 6.42, 11.5, 0.3,
         [[B("Targets and timing will be refined through the pilot. ", color=NAVY), ("These are ambitions, not current traction.", {"color": MUTED})]], size=11)
    footer(s, 18)
    notes(s, "We are in the NOW phase: interviewing founders, gathering contributions, choosing which services to test and defining platform requirements. "
             "NEXT is a deliberately small pilot where we measure usefulness and delivery cost. THEN we repeat what works. "
             "The milestones are targets: 10,000 free users is the threshold for a public launch event, March 2027 is a tentative target subject to readiness, and 100,000 is the wider ambition. None is current traction.")


def s_people():
    global s
    s = blank(prs)
    header(s, 19, "People and collaboration", "Bringing complementary experience together.")
    clusters = [("Strategy and commercial", ["Commercial leadership", "GTM and customer lifecycle", "Legal and governance"], NAVY),
                ("Founder support", ["Startup and incubation experience", "Consumer-business mentoring", "Education and university outreach"], BLUE),
                ("Platform and growth", ["AI and business systems", "Branding and growth"], "1A55AC"),
                ("Community and delivery", ["Operations and community", "Events and experiences"], "17407F")]
    cxs = [ML, ML + 2.05 * 1, ML, ML]
    qx, qy, qw, qh = ML, 2.2, 3.7, 2.2
    for i, (cl, caps, col) in enumerate(clusters):
        col_i, row_i = i % 2, i // 2
        x = qx + col_i * (qw + 0.15)
        y = qy + row_i * (qh + 0.15)
        rect(s, x, y, qw, qh, fill=PALE)
        rect(s, x, y, qw, 0.07, fill=col)
        text(s, x + 0.22, y + 0.22, qw - 0.4, 0.3, cl.upper(), size=9.5, color=col, bold=True, spacing=1.1)
        for k, cp in enumerate(caps):
            cc_ = rect(s, x + 0.22, y + 0.62 + k * 0.5, qw - 0.44, 0.4, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            shape_text(cc_, cp, size=11, color=NAVY, bold=True, align=PP_ALIGN.LEFT, margin=0.16)
    rx = 8.5
    text(s, rx, 2.2, 4.18, 1.6, "We are exploring how like-minded contributors can participate through their experience, resources and available time.",
         size=15, color=NAVY, bold=True, line=1.25)
    text(s, rx, 3.75, 4.5, 0.3, "HOW ROLES WILL TAKE SHAPE", size=9.5, color=BLUE, bold=True, spacing=1.3)
    text(s, rx, 4.1, 4.18, 1.4,
         "Longer-term roles and responsibilities will be considered as the project develops, based on contribution, time commitment, responsibility taken, reliable delivery and value created.",
         size=12, color=TEXT, line=1.3)
    line(s, rx, 5.65, W - ML, 5.65, color=LINE)
    text(s, rx, 5.8, 4.5, 0.7, [[B("Ritesh Kumar", color=NAVY)], [("Project Initiator", {"color": MUTED, "size": 12})]], size=14, line=1.2)
    footer(s, 19)
    notes(s, "This is a capability map, not an org chart. It shows the mix of experience the project is bringing together across commercial, founder support, platform and community delivery. "
             "Contributors are participating during the shaping and pilot phase. Permanent roles, titles and any longer-term arrangements have not been set and will be considered later based on actual contribution and delivery.")


def s_invitation():
    global s
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.18, H, fill=BLUE)
    text(s, ML + 0.1, 0.48, 8, 0.3, [[(eyebrow_num(20) + "INVITATION TO PARTICIPATE", {"bold": True})]], size=10, color=MID, spacing=1.5)
    text(s, ML + 0.1, 0.82, 8.6, 1.3, "Help shape a practical ecosystem for the founder journey.", size=30, color=WHITE, bold=True, line=1.05)
    ctas = [("Founders", "Help shape the platform and express interest in the pilot."),
            ("Mentors and specialists", "Contribute expertise and practical sessions."),
            ("Partners and universities", "Explore a useful collaboration."),
            ("Investors", "Connect around founder readiness and ecosystem opportunities."),
            ("Prospective collaborators", "Discuss how your skills and time can contribute.")]
    for i, (a, b) in enumerate(ctas):
        y = 2.35 + i * 0.82
        line(s, ML + 0.1, y, 8.6, y, color="2E4A72", w=0.75)
        text(s, ML + 0.1, y + 0.17, 2.9, 0.5, a, size=14.5, color=WHITE, bold=True)
        text(s, 3.85, y + 0.17, 4.8, 0.55, b, size=13, color="C9D6E8", line=1.2)
    qr = rect(s, 9.55, 2.3, 3.1, 3.1, fill=WHITE)
    s.shapes.add_picture(qr_png(URL), Inches(9.7), Inches(2.45), Inches(2.8), Inches(2.8))
    text(s, 9.55, 5.55, 3.1, 0.35, "startuproof.com", size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, 9.55, 5.95, 3.1, 0.5, "Scan to visit the website and get in touch.", size=11, color="9FB2CC", align=PP_ALIGN.CENTER)
    wordmark(s, ML + 0.1, 6.85, size=11, dark=True, w=4)
    text(s, 5.0, 6.87, W - ML - 5.0, 0.3, "An ecosystem around the founder. A digital workspace for the journey.", size=10.5, color="9FB2CC", align=PP_ALIGN.RIGHT)
    notes(s, "Close with a specific ask for each person in the room. Founders: help shape the pilot. Mentors and specialists: offer a session. "
             "Partners and universities: explore a first, small collaboration. Investors: engage around founder readiness. Collaborators: tell us what skills and time you can offer. "
             "Point to the website and QR code as the way to get in touch, and invite a follow-up conversation.")



SHORT_ORDER = [s_cover, s_problem, s_what, s_who, s_journey, s_layers, s_founder_os, s_scenario, s_ai_human,
               s_community, s_academy, s_services, s_insights, s_incubation, s_investors, s_partners,
               s_commercial, s_roadmap, s_people, s_invitation]

if __name__ == "__main__":
    OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "Startup_Roof_General_Presentation.pptx")
    prs = new_prs()
    for fn in SHORT_ORDER:
        fn()
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")